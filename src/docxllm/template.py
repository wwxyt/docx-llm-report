"""生成示例 Word 模版（程序化写入 Jinja2 标签）。

模版标签的三条规则：
- 单值用 {{ var }}；
- 表格行循环用 {%tr for r in rows %} + {%tr endfor %}，且两个标签必须各占一整行；
- 段落级循环用 {%p for p in analysis %} + {%p endfor %}，中间放 {{ p }}。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

FONT = "Microsoft YaHei"


def _style(run, size=10.5, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _para(doc, text, size=10.5, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    _style(p.add_run(text), size=size, bold=bold)
    return p


def build_template(
    dest: str | Path,
    rows_var: str = "rows",
    table_fields: tuple[str, str, str] = ("region", "sales", "yoy"),
    table_headers: tuple[str, str, str] = ("地区", "销售额（万元）", "同比"),
) -> Path:
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    _para(doc, "{{ title }}", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _para(doc, "报告期间：{{ period }}", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    _para(doc, "一、总体情况", size=14, bold=True, space_after=6)
    _para(doc, "{{ summary }}", space_after=12)

    _para(doc, "二、关键指标", size=14, bold=True, space_after=6)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(table_headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        _style(cell.paragraphs[0].add_run(h), bold=True)

    table.rows[1].cells[0].text = "{%%tr for r in %s %%}" % rows_var
    body = table.rows[2].cells
    for i, field in enumerate(table_fields):
        body[i].text = "{{ r.%s }}" % field
    table.rows[3].cells[0].text = "{%tr endfor %}"
    for row in (table.rows[1], table.rows[2], table.rows[3]):
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    _style(run)

    _para(doc, "三、趋势分析", size=14, bold=True, space_after=6)
    _para(doc, "{%p for p in analysis %}", space_after=0)
    _para(doc, "{{ p }}")
    _para(doc, "{%p endfor %}", space_after=12)

    _para(doc, "四、结论与建议", size=14, bold=True, space_after=6)
    _para(doc, "{{ conclusion }}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(24)
    _style(p.add_run("生成时间：{{ generated_at }}"), size=9)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style(p2.add_run("{{ model_note }}"), size=8)

    doc.save(dest)
    return dest

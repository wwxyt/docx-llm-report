"""Excel 读取、字段映射与数据摘要。

设计原则：数字只从数据源来，绝不交给模型去算。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .config import DataSpec


def read_sheet(path: str | Path, sheet: str | None = None) -> tuple[list[str], list[dict]]:
    """读取首个（或指定）工作表，返回 (表头, 行字典列表)。"""
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path} 是空表")
    headers = [str(h).strip() if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
    data: list[dict] = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        data.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
    if not data:
        raise ValueError(f"{path} 除表头外没有数据行")
    return headers, data


def load_data(spec: DataSpec, base_path: str | Path) -> tuple[list[str], list[dict], list[dict]]:
    """按配置读取并映射字段，返回 (表头, 原始行, 模版可用行)。"""
    headers, rows = read_sheet(base_path, spec.sheet)
    if spec.columns:
        missing = [c for c in spec.columns if c not in headers]
        if missing:
            raise ValueError(f"Excel 中找不到配置的表头：{missing}；实际表头：{headers}")
        mapped = [{var: row.get(col) for col, var in spec.columns.items()} for row in rows]
    else:
        mapped = rows
    return headers, rows, mapped


def to_number(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("%", "").replace("+", "")
    try:
        return float(text)
    except ValueError:
        return None


def summarize(headers: list[str], rows: list[dict]) -> str:
    """为模型提供一份"算好"的数值摘要，避免它自己算错。"""
    lines: list[str] = [f"数据共 {len(rows)} 行，字段：{'、'.join(headers)}。"]
    for h in headers:
        values = [to_number(r.get(h)) for r in rows]
        nums = [v for v in values if v is not None]
        if len(nums) < 2 or len(nums) < len(values) * 0.8:
            continue
        total = sum(nums)
        idx_max = max(range(len(values)), key=lambda i: values[i] if values[i] is not None else float("-inf"))
        idx_min = min(range(len(values)), key=lambda i: values[i] if values[i] is not None else float("inf"))
        first_header = headers[0]
        lines.append(
            f"{h}：合计 {total:,.2f}，均值 {total / len(nums):,.2f}，"
            f"最大 {nums and max(nums):,.2f}（{rows[idx_max].get(first_header)}），"
            f"最小 {min(nums):,.2f}（{rows[idx_min].get(first_header)}）"
        )
    return "\n".join(lines)


def to_markdown(headers: list[str], rows: list[dict]) -> str:
    lines = ["| " + " | ".join(headers) + " |",
             "| " + " | ".join("---" for _ in headers) + " |"]
    for r in rows:
        lines.append("| " + " | ".join(str(r.get(h, "")) for h in headers) + " |")
    return "\n".join(lines)


def write_sample(path: str | Path) -> Path:
    """生成示例数据，供 init 与 demo 使用。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "销售数据"
    ws.append(["地区", "销售额（万元）", "同比"])
    for row in [
        ("华东", 3280.5, "+12.4%"),
        ("华南", 2140.0, "+8.1%"),
        ("华北", 1875.3, "-2.6%"),
        ("西南", 960.8, "+21.7%"),
        ("东北", 612.4, "-5.3%"),
    ]:
        ws.append(row)
    wb.save(path)
    return path

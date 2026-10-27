"""端到端测试：全部离线运行，不触网、不需要 API Key。"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from docxllm.cli import SAMPLE_CONFIG, main
from docxllm.config import load_config
from docxllm.excel import write_sample
from docxllm.llm import parse_json
from docxllm.pipeline import generate
from docxllm.template import build_template
from docxllm.validate import check, extract_usage


@pytest.fixture()
def project(tmp_path: Path) -> Path:
    write_sample(tmp_path / "examples" / "sample.xlsx")
    build_template(tmp_path / "templates" / "report.docx")
    (tmp_path / "report.yaml").write_text(SAMPLE_CONFIG, encoding="utf-8")
    return tmp_path


def _all_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(c.text for c in row.cells)
    return "\n".join(chunks)


def test_mock_pipeline_generates_docx(project: Path):
    cfg = load_config(project / "report.yaml")
    out, content = generate(cfg, out=project / "output" / "r.docx", mock=True)

    assert out.exists()
    assert content["title"] == "示例分析报告"
    assert len(content["analysis"]) == 3

    text = _all_text(out)
    assert "示例分析报告" in text
    assert "一、总体情况" in text
    assert "四、结论与建议" in text
    # 模版标签必须全部被替换
    assert "{{" not in text and "{%" not in text


def test_table_rows_follow_excel(project: Path):
    cfg = load_config(project / "report.yaml")
    out, _ = generate(cfg, out=project / "output" / "r.docx", mock=True)
    table = Document(out).tables[0]

    assert table.rows[0].cells[0].text == "地区"
    assert len(table.rows) == 6  # 1 表头 + 5 数据行
    assert [c.text for c in table.rows[1].cells] == ["华东", "3280.5", "+12.4%"]
    assert [c.text for c in table.rows[5].cells] == ["东北", "612.4", "-5.3%"]


def test_analysis_paragraphs_expand(project: Path):
    cfg = load_config(project / "report.yaml")
    out, _ = generate(cfg, out=project / "output" / "r.docx", mock=True)
    paras = [p.text for p in Document(out).paragraphs]
    hits = [p for p in paras if p.startswith("（mock 第")]
    assert len(hits) == 3
    assert any("mock 第 3 条" in p for p in hits)


def test_columns_mapping_applies(project: Path):
    cfg = load_config(project / "report.yaml")
    cfg.data.columns = {"地区": "region", "销售额（万元）": "sales"}
    out, _ = generate(cfg, out=project / "output" / "r.docx", mock=True)
    assert Document(out).tables[0].rows[1].cells[2].text == ""


def test_missing_column_raises(project: Path):
    cfg = load_config(project / "report.yaml")
    cfg.data.columns = {"不存在的表头": "x"}
    with pytest.raises(ValueError, match="找不到配置的表头"):
        generate(cfg, out=project / "output" / "r.docx", mock=True)


def test_parse_json_handles_fence_and_noise():
    assert parse_json('{"a": 1}') == {"a": 1}
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('好的，结果如下：\n{"a": 1}\n以上。') == {"a": 1}
    assert parse_json('```\n{"a": 1}\n```') == {"a": 1}


def test_check_passes_on_sample(project: Path):
    errors, _ = check(load_config(project / "report.yaml"))
    assert errors == []


def test_check_reports_unknown_variable(project: Path):
    cfg = load_config(project / "report.yaml")
    cfg.llm.sections[1].var = "sumary"  # 模版里仍是 summary
    errors, warnings = check(cfg)
    assert any("summary" in e for e in errors)
    assert any("sumary" in w for w in warnings)


def test_extract_usage_finds_loops(project: Path):
    usage = extract_usage(project / "templates" / "report.docx")
    assert "rows" in usage["loop_sources"]
    assert "analysis" in usage["loop_sources"]
    assert "p" in usage["loop_vars"] and "r" in usage["loop_vars"]


def test_cli_init_and_run(project: Path, capsys):
    target = project / "demo"
    assert main(["init", "--dir", str(target)]) == 0
    assert (target / "report.yaml").exists()
    assert main(["run", "-c", str(target / "report.yaml"), "--mock",
                 "-o", str(target / "output" / "x.docx")]) == 0
    assert (target / "output" / "x.docx").exists()
    assert main(["check", "-c", str(target / "report.yaml")]) == 0
